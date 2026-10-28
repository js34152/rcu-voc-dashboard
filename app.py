# -*- coding: utf-8 -*-
from collections import Counter
from datetime import datetime
import io
import json
import os
import re
import urllib.request
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

import importlib
import translations as tr
importlib.reload(tr)

# -------------------------------------------------------------
# 0. 기본 설정 & 사내 보안 게이트 (암호: rcu2026!)
# -------------------------------------------------------------
st.set_page_config(
    page_title="Global RCU VOC Intelligence v3.0",
    page_icon="📡",
    layout="wide",
    initial_sidebar_state="expanded",
)

LANG_OPTIONS = {
    "🇰🇷 한국어 (Korean)": "ko",
    "🇺🇸 English": "en",
    "🇨🇳 简体中文 (Chinese)": "zh",
    "🇻🇳 Tiếng Việt (Vietnamese)": "vi",
    "🇹🇭 ภาษาไทย (Thai)": "th",
}

if "lang" not in st.session_state:
    st.session_state["lang"] = "ko"

# 사이드바 최상단 언어 선택 (로그인 전/후 공통 작동)
selected_lang_label = st.sidebar.selectbox(
    "🌐 Language / 언어 선택",
    options=list(LANG_OPTIONS.keys()),
    index=list(LANG_OPTIONS.values()).index(st.session_state["lang"]),
)
curr_lang = LANG_OPTIONS[selected_lang_label]
st.session_state["lang"] = curr_lang


def txt(key, **kwargs):
    s = tr.UI_TEXT.get(curr_lang, {}).get(key, tr.UI_TEXT["ko"].get(key, key))
    if kwargs:
        return s.format(**kwargs)
    return s


def get_cat_name(k):
    return tr.CATEGORIES.get(k, {}).get(curr_lang, k)


def get_dev_name(k):
    return tr.DEVICE_TYPES.get(k, {}).get(curr_lang, k)


def get_power_name(k):
    return tr.POWER_TYPES.get(k, {}).get(curr_lang, k)


def get_ui_name(k):
    return tr.UI_TYPES.get(k, {}).get(curr_lang, k)


def get_region_name(k):
    return tr.REGIONS.get(k, {}).get(curr_lang, k)


def get_col_name(k):
    return tr.TABLE_COLUMNS.get(k, {}).get(curr_lang, k)


def get_sentiment_name(k):
    return tr.SENTIMENTS.get(k, {}).get(curr_lang, k)


def get_action_sol(k):
    return tr.ACTION_RECOMMENDATIONS.get(k, {}).get(
        curr_lang, tr.ACTION_RECOMMENDATIONS.get(k, {}).get("ko", "")
    )


def get_secret_api_key():
    try:
        return st.secrets.get("GEMINI_API_KEY", os.environ.get("GEMINI_API_KEY", ""))
    except Exception:
        return os.environ.get("GEMINI_API_KEY", "")


RCU_PASSWORD = "rcu2026!"


def check_password():
    if st.session_state.get("password_correct", False):
        return True

    def password_entered():
        if st.session_state.get("password_input") == RCU_PASSWORD:
            st.session_state["password_correct"] = True
            del st.session_state["password_input"]
        else:
            st.session_state["password_correct"] = False

    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.title(txt("gate_title"))
        st.markdown(f"##### {txt('gate_sub')}")
        st.info(txt("gate_info"))

        st.text_input(
            txt("gate_input"),
            type="password",
            on_change=password_entered,
            key="password_input",
            placeholder=txt("gate_placeholder"),
        )

        if "password_correct" in st.session_state and not st.session_state.get(
            "password_correct"
        ):
            st.error(txt("gate_error"))

    return False


if not check_password():
    st.stop()

if st.sidebar.button(txt("logout"), help=txt("logout_help")):
    st.session_state["password_correct"] = False
    st.rerun()

# -------------------------------------------------------------
# 1. 60개 브랜드 디렉터리 & 폼팩터 매핑
# -------------------------------------------------------------
BRAND_DIRECTORY = {
    # TV 제조사
    "Samsung": {"region": "Korea", "type": "TV 제조사", "power": "솔라셀/C타입 충전식", "ui": "미니멀 물리키"},
    "LG Electronics": {"region": "Korea", "type": "TV 제조사", "power": "건전지 교체형", "ui": "자이로/에어마우스"},
    "Sony": {"region": "Japan", "type": "TV 제조사", "power": "건전지 교체형", "ui": "백라이트 물리키"},
    "Panasonic": {"region": "Japan", "type": "TV 제조사", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "Sharp": {"region": "Japan", "type": "TV 제조사", "power": "건전지 교체형", "ui": "다버튼 핫키형"},
    "TCL": {"region": "Global", "type": "TV 제조사", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Hisense": {"region": "Global", "type": "TV 제조사", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Xiaomi": {"region": "ASIA", "type": "TV 제조사", "power": "건전지 교체형", "ui": "초미니멀 물리키"},
    "Vizio": {"region": "NA", "type": "TV 제조사", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Philips": {"region": "EU", "type": "TV 제조사", "power": "건전지 교체형", "ui": "양면/특수 물리키"},
    # 스트리밍 / 스마트 미디어
    "Apple": {"region": "Global", "type": "스트리밍/스마트 미디어", "power": "C타입 충전식", "ui": "터치 클릭패드"},
    "Amazon Fire TV": {"region": "Global", "type": "스트리밍/스마트 미디어", "power": "건전지 교체형", "ui": "백라이트 물리키"},
    "Google": {"region": "Global", "type": "스트리밍/스마트 미디어", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Roku": {"region": "NA", "type": "스트리밍/스마트 미디어", "power": "C타입 충전식", "ui": "이어폰잭/물리키"},
    "NVIDIA": {"region": "Global", "type": "스트리밍/스마트 미디어", "power": "건전지 교체형", "ui": "삼각기둥 백라이트"},
    "TiVo": {"region": "NA", "type": "스트리밍/스마트 미디어", "power": "건전지 교체형", "ui": "인체공학 피넛키"},
    "SofaBaton": {"region": "Global", "type": "스트리밍/스마트 미디어", "power": "C타입 충전식", "ui": "OLED스크린/휠"},
    # 방송통신사 (Telco/MSO)
    "Comcast": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "음성 특화 물리키"},
    "Charter": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "DirecTV": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "안드로이드 물리키"},
    "DISH Network": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "백라이트 물리키"},
    "Cox": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "음성 특화 물리키"},
    "Rogers": {"region": "NA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "음성 특화 물리키"},
    "Sky Group": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "터치/모션 백라이트"},
    "Vodafone": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "LGI (Virgin Media)": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "음성 특화 물리키"},
    "Canal+": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Orange": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Deutsche Telekom": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Telefonica": {"region": "EU", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "KT (지니TV)": {"region": "Korea", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "SKB (B tv)": {"region": "Korea", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "이어폰일체형 물리키"},
    "LGU+ (U+tv)": {"region": "Korea", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "이어폰일체형 물리키"},
    "MultiChoice (DStv)": {"region": "Africa", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "StarTimes": {"region": "Africa", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Canal+ Afrique": {"region": "Africa", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Tata Play": {"region": "ASIA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 핫키형"},
    "Airtel": {"region": "ASIA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "J:COM": {"region": "ASIA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "Hikari TV": {"region": "ASIA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 하이브리드"},
    "Foxtel": {"region": "ASIA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Claro": {"region": "SCA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    "Vrio Corp": {"region": "SCA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "다버튼 물리키"},
    "Televisa (izzi)": {"region": "SCA", "type": "방송통신사 (Telco/MSO)", "power": "건전지 교체형", "ui": "미니멀 물리키"},
    # 스마트홈 오토메이션
    "Control4": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "터치스크린/물리키"},
    "Crestron": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "터치스크린/물리키"},
    "Savant": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "고해상도 터치스크린"},
    "Josh.ai": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "AI음성/물리키"},
    "URC": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "터치스크린/물리키"},
    "RTI": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "대화면 터치스크린"},
    "Nice (ELAN)": {"region": "NA", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "터치스크린/인터콤"},
    "Basalte": {"region": "EU", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "명품 금속 터치"},
    "Logitech": {"region": "Global", "type": "스마트홈 오토메이션", "power": "크래들 무선충전식", "ui": "터치스크린/물리키"},
    "ADT": {"region": "NA", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "보안 키패드"},
    "Vivint": {"region": "NA", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "스마트 패널"},
    "Resideo": {"region": "NA", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "보안 키패드"},
    "Loxone": {"region": "EU", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "5터치 포인트"},
    "Somfy": {"region": "EU", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "모터 제어 휠"},
    "Cync (GE)": {"region": "NA", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "조명 디밍 버튼"},
    "Tuya": {"region": "Global", "type": "스마트홈 오토메이션", "power": "건전지 교체형", "ui": "스마트 허브 연동"},
}

BRAND_ALIASES = {
    "amazon": "Amazon Fire TV", "fire tv": "Amazon Fire TV", "google tv": "Google",
    "comcast (xfinity)": "Comcast", "xfinity": "Comcast", "spectrum": "Charter",
    "sky": "Sky Group", "virgin media": "LGI (Virgin Media)", "canal": "Canal+",
    "dt": "Deutsche Telekom", "telekom": "Deutsche Telekom", "vodafone tv": "Vodafone",
    "dstv": "MultiChoice (DStv)", "startimes": "StarTimes", "b tv": "SKB (B tv)",
    "skb": "SKB (B tv)", "지니tv": "KT (지니TV)", "kt": "KT (지니TV)", "u+tv": "LGU+ (U+tv)",
    "lgu+": "LGU+ (U+tv)", "lg": "LG Electronics", "samsung": "Samsung",
}


def normalize_brand(name):
    raw = str(name).strip()
    if raw in BRAND_DIRECTORY:
        return raw
    lower = raw.lower()
    if lower in BRAND_ALIASES:
        return BRAND_ALIASES[lower]
    for alias, canon in BRAND_ALIASES.items():
        if alias in lower:
            return canon
    for canon in BRAND_DIRECTORY:
        if canon.lower() in lower:
            return canon
    return raw

# -------------------------------------------------------------
# 2. 지역별 Champions 매트릭스
# -------------------------------------------------------------
REGIONAL_TOP3 = {
    "한국 (KR)": {
        "TV 제조사": ["Samsung", "LG Electronics"],
        "방송통신사 (Telco/MSO)": ["KT (지니TV)", "SKB (B tv)", "LGU+ (U+tv)"],
    },
    "북미 (NA)": {
        "TV 제조사": ["Samsung", "Vizio", "TCL"],
        "스트리밍/스마트 미디어": ["Roku", "Amazon Fire TV", "Apple"],
        "방송통신사 (Telco/MSO)": ["Comcast", "Charter", "DirecTV"],
        "스마트홈 오토메이션": ["Control4", "Crestron", "Savant"],
    },
    "유럽 (EU)": {
        "TV 제조사": ["Samsung", "Philips", "LG Electronics"],
        "스트리밍/스마트 미디어": ["Amazon Fire TV", "Apple", "Google"],
        "방송통신사 (Telco/MSO)": ["Sky Group", "Vodafone", "Orange"],
        "스마트홈 오토메이션": ["Somfy", "Loxone", "Basalte"],
    },
    "아시아-태평양 (APAC)": {
        "TV 제조사": ["Sony", "Panasonic", "Sharp"],
        "스트리밍/스마트 미디어": ["Google", "Xiaomi", "Apple"],
        "방송통신사 (Telco/MSO)": ["Tata Play", "Airtel", "J:COM"],
        "스마트홈 오토메이션": ["Tuya", "SofaBaton", "Logitech"],
    },
    "아프리카 (Africa)": {
        "TV 제조사": ["Samsung", "Hisense", "LG Electronics"],
        "방송통신사 (Telco/MSO)": ["MultiChoice (DStv)", "StarTimes", "Canal+ Afrique"],
    },
}

KR_CHAMPIONS = sum(REGIONAL_TOP3["한국 (KR)"].values(), [])
NA_CHAMPIONS = sum(REGIONAL_TOP3["북미 (NA)"].values(), [])
EU_CHAMPIONS = sum(REGIONAL_TOP3["유럽 (EU)"].values(), [])
APAC_CHAMPIONS = sum(REGIONAL_TOP3["아시아-태평양 (APAC)"].values(), [])
AFRICA_CHAMPIONS = sum(REGIONAL_TOP3["아프리카 (Africa)"].values(), [])


def refine_category(text):
    t = str(text).lower()
    if any(
        w in t
        for w in [
            "pair", "connect", "bluetooth", "wifi", "offline", "drop",
            "disconnect", "sync", "unpair", "network", "link", "lost",
            "load shedding", "synchro", "verbindung", "페어링", "연결",
        ]
    ):
        return "페어링/연결 끊김"
    if any(
        w in t
        for w in [
            "lag", "slow", "freeze", "stuck", "delay", "unresponsive",
            "hang", "epg", "verzögerung", "langsam", "지연", "렉",
            "먹통", "느려", "멈춤",
        ]
    ):
        return "속도/지연/먹통"
    if any(
        w in t
        for w in [
            "volume", "mute", "cec", "hdmi", "turn on", "turn off",
            "soundbar", "receiver", "tv", "power", "handshake", "lautstärke",
            "볼륨", "소리", "전원", "외부입력",
        ]
    ):
        return "볼륨/TV연동(CEC)"
    if any(
        w in t
        for w in [
            "range", "line of sight", "distance", "angle", "ir ", "signal",
            "sensor", "aim", "reichweite", "수신", "거리", "각도", "정면",
        ]
    ):
        return "신호범위/IR센서"
    if any(
        w in t
        for w in [
            "button", "layout", "press", "key", "touch", "click", "backlight",
            "ui", "screen", "dust", "sand", "stiff", "taste", "drücken",
            "bouton", "버튼", "키", "터치", "클릭", "더블클릭", "채터링", "백라이트",
        ]
    ):
        return "버튼/UI 제어 불량"
    if any(
        w in t
        for w in [
            "battery", "power", "charge", "drain", "solar", "recharge",
            "corrosion", "leak", "batterie", "akku", "pile", "배터리",
            "건전지", "충전", "방전", "광탈", "누액", "부식",
        ]
    ):
        return "배터리/전원"
    if any(
        w in t
        for w in [
            "voice", "mic", "alexa", "bixby", "siri", "assistant",
            "voix", "stimme", "음성", "마이크", "인식", "기가지니", "클로바",
        ]
    ):
        return "음성 인식/마이크"
    if any(
        w in t
        for w in [
            "cheap", "plastic", "break", "broke", "drop", "fall",
            "durability", "cassé", "kaputt", "깨짐", "파손", "마감", "부러", "내구성",
        ]
    ):
        return "내구성/하드웨어 퀄리티"
    if any(
        w in t
        for w in [
            "login", "sign in", "account", "update", "ad ", "mise à jour",
            "업데이트", "계정", "광고",
        ]
    ):
        return "계정/업데이트/광고"
    if any(
        w in t
        for w in [
            "app", "hotkey", "shortcut", "finder", "raccourci", "핫키",
            "바로가기", "찾기", "비프음",
        ]
    ):
        return "단순 앱/기능 불만"
    return "기타/일반"


NOISE_FILTER_REGEX = re.compile(
    r"(?:공식\s*(?:기업\s*)?사이트|공식\s*홈페이지|공식\s*(?:직영\s*)?온라인샵|다이렉트샵|"
    r"\[kt\]\s*메인|\[마이페이지\]|고객센터|음성상담|채팅상담|보이스봇|"
    r"이용하시는\s*고객|혜택과\s*서비스|이벤트도\s*참여|되겠습니다|가능합니다|"
    r"주식회사|사업자등록|대표이사|개인정보처리|요금제\s*확인|상품서비스|"
    r"delivery|shipping|hassle-free|returns?|cash on delivery|"
    r"shop for|buy online|in stock|order now|price in|naira|rand|"
    r"nationwide delivery|collection|many ways to pay|money back|"
    r"is designed for|designed to provide|high-quality replacement|"
    r"seamless compatibility|official website|features an? |leading provider|"
    r"overview|specifications?:|package includes?:|replacement remote|aerial king)",
    re.IGNORECASE,
)


@st.cache_data(ttl=300)
def load_accumulated_data():
    if not os.path.exists("voc_db.csv"):
        return pd.DataFrame()

    try:
        df = pd.read_csv("voc_db.csv", encoding="utf-8-sig")
    except Exception:
        df = pd.read_csv("voc_db.csv", encoding="utf-8")

    if not df.empty:
        df = df[~df["Review"].astype(str).str.contains(NOISE_FILTER_REGEX)].copy()
        df = df[df["Review"].astype(str).str.len() >= 25].copy()

        df["Brand"] = df["Brand"].apply(normalize_brand)
        df["Device_Type"] = df["Brand"].map(
            lambda b: BRAND_DIRECTORY.get(b, {}).get("type", "TV 제조사")
        )
        df["Region"] = df["Brand"].map(
            lambda b: BRAND_DIRECTORY.get(b, {}).get("region", "Global")
        )
        df["Power_Type"] = df["Brand"].map(
            lambda b: BRAND_DIRECTORY.get(b, {}).get("power", "건전지 교체형")
        )
        df["UI_Type"] = df["Brand"].map(
            lambda b: BRAND_DIRECTORY.get(b, {}).get("ui", "미니멀 물리키")
        )

        df["Category"] = df["Review"].apply(refine_category).fillna("기타/일반")
        df["Sentiment"] = df["Sentiment"].fillna("Negative")
        if "Source" not in df.columns:
            df["Source"] = "Web / Community"

        if "Collected_Date" in df.columns:
            df["Parsed_Date"] = pd.to_datetime(
                df["Collected_Date"].str.slice(0, 10), errors="coerce"
            )
            df["Parsed_Date"] = df["Parsed_Date"].fillna(pd.Timestamp.now())
        else:
            df["Parsed_Date"] = pd.Timestamp.now()

    return df


df = load_accumulated_data()
if df.empty:
    st.warning("⚠️ 'voc_db.csv' is empty or missing.")
    st.stop()


# -------------------------------------------------------------
# Gemini AI 연동 엔진 (3줄 요약 & Root Cause 진단)
# -------------------------------------------------------------
def analyze_voc_with_gemini(api_key, client_name, p_issue, s_issue, sample_reviews, lang="ko"):
    api_key = str(api_key).strip().strip('"').strip("'")
    if not api_key:
        return "⚠️ Gemini API Key가 입력되지 않았습니다."

    lang_names = {
        "ko": "Korean (한국어)",
        "en": "English",
        "zh": "Simplified Chinese (简体中文)",
        "vi": "Vietnamese (Tiếng Việt)",
        "th": "Thai (ภาษาไทย)"
    }
    target_lang_name = lang_names.get(lang, "Korean")
    reviews_formatted = "\n".join([f"- {r}" for r in sample_reviews])

    sum_title = tr.UI_TEXT.get(lang, {}).get("ai_summary_title", "고객 VoC 핵심 3줄 요약")
    rc_title = tr.UI_TEXT.get(lang, {}).get("ai_root_cause_title", "H/W & F/W 결함 근본 원인 (Root Cause) 분석")
    st_title = tr.UI_TEXT.get(lang, {}).get("ai_strategy_title", "B2B 수주전 기술 차별화 제언")

    prompt = f"""You are a distinguished Principal Remote Control Unit (RCU) Hardware/Firmware Architect and Telco B2B Solution Leader.
Analyze the following Voice of Customer (VoC) feedback for target client '{client_name}'.

- 1st Priority Defect: {p_issue}
- 2nd Priority Defect: {s_issue if s_issue else "None"}
- Field Customer Verbatim Reviews:
{reviews_formatted}

Generate an executive technical diagnosis strictly in {target_lang_name}.
Format the output into exactly three sections using Markdown:

### 📋 {sum_title}
(Provide 3 concise, impactful bullet points summarizing the actual user pain points.)

### 🔬 {rc_title}
(Provide 2-3 deep engineering root-cause diagnoses explaining the hardware/firmware bottlenecks, such as MCU wake-up jitter, BLE advertising/connection intervals, CEC handshake bus collisions, power management IC leakage, or metal dome debounce circuit timing.)

### 🎯 {st_title}
(Provide 2 concrete technical differentiation proposals to decisively win the B2B supply contract over competitors.)

Be concise, highly professional, and engineering-grade. Do not include fluff."""

    non_text_keywords = ["tts", "audio", "embed", "image", "imagen", "robotics", "realtime", "live"]
    active_models = []
    try:
        list_url = f"https://generativelanguage.googleapis.com/v1beta/models?key={api_key}"
        req = urllib.request.Request(list_url, headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=12) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            for m in data.get("models", []):
                m_name = m.get("name", "")
                if "generateContent" in m.get("supportedGenerationMethods", []):
                    if not any(bad in m_name.lower() for bad in non_text_keywords):
                        active_models.append(m_name)
    except urllib.error.HTTPError as e:
        try:
            err_data = json.loads(e.read().decode("utf-8"))
            err_msg = err_data.get("error", {}).get("message", str(e))
        except Exception:
            err_msg = str(e)
        if e.code in (401, 403):
            return f"⚠️ Gemini API 인증/키 오류 ({err_msg})"
    except Exception:
        pass

    preferred_order = [
        "models/gemini-2.0-flash",
        "models/gemini-1.5-flash",
        "models/gemini-1.5-flash-latest",
        "models/gemini-1.5-pro",
        "models/gemini-pro"
    ]
    
    candidate_models = []
    for pref in preferred_order:
        if pref in active_models:
            candidate_models.append(pref)
    for m in active_models:
        if m not in candidate_models:
            candidate_models.append(m)

    if not candidate_models:
        candidate_models = preferred_order

    last_err = ""
    for model_full_name in candidate_models:
        clean_name = model_full_name if model_full_name.startswith("models/") else f"models/{model_full_name}"
        url = f"https://generativelanguage.googleapis.com/v1beta/{clean_name}:generateContent?key={api_key}"
        headers = {"Content-Type": "application/json"}
        payload = {
            "contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"temperature": 0.3, "maxOutputTokens": 1500}
        }
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers)
            with urllib.request.urlopen(req, timeout=35) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                return data["candidates"][0]["content"]["parts"][0]["text"]
        except urllib.error.HTTPError as e:
            try:
                err_data = json.loads(e.read().decode("utf-8"))
                msg = err_data.get("error", {}).get("message", str(e))
            except Exception:
                msg = str(e)
            last_err = f"HTTP {e.code}: {msg}"
            if e.code in (401, 403):
                return f"⚠️ Gemini API 인증 오류 ({last_err})"
            continue
        except Exception as e:
            last_err = str(e)
            continue

    return f"⚠️ Gemini API 분석 실패: {last_err}"


# -------------------------------------------------------------
# PPTX 다국어 생성 엔진 (네이티브 벡터 차트 & AI 슬라이드 탑재)
# -------------------------------------------------------------
def create_pptx_deck(
    client_name, p_issue, s_issue, p_sol, s_sol, pitch, reviews, cat_counts=None, ai_analysis=None, lang="ko"
):
    prs = Presentation()
    t_titles = {
        "ko": {
            "deck_title": f"[B2B Pitch 제안서]\n{client_name} RCU 품질 혁신 & UX 개선안",
            "date": "작성일자", "target": "대상 고객사",
            "s1_title": "1. 고객 핵심 취약점 (VoC Pain Points)",
            "p_rank1": "🚨 1순위 결함 지수:", "p_rank2": "⚠️ 2순위 결함 지수:",
            "reviews_header": "💬 현장 고객 대표 불만 원문:",
            "s2_title": "2. 당사 R&D 기술 해결 솔루션", "sol_for": "대응 솔루션:",
            "s3_title": "3. 임원 보고 / B2B 수주 맞춤 Pitch",
            "exec_pitch": f"Executive Pitch for {client_name}",
            "chart_series": "불만 점유율", "others": "기타",
            "ai_slide_title": "4. Gemini AI 심층 결함 원인 (Root Cause) 진단",
        },
        "en": {
            "deck_title": f"[B2B Pitch Proposal]\n{client_name} RCU Quality & UX Improvement",
            "date": "Date", "target": "Target Brand",
            "s1_title": "1. Core Customer Pain Points (VoC)",
            "p_rank1": "🚨 1st Priority Defect Index:", "p_rank2": "⚠️ 2nd Priority Defect Index:",
            "reviews_header": "💬 Representative Verbatim Complaints:",
            "s2_title": "2. Proprietary R&D Engineering Countermeasures", "sol_for": "Technical Countermeasure:",
            "s3_title": "3. Executive / B2B Proposal Pitch",
            "exec_pitch": f"Executive Pitch for {client_name}",
            "chart_series": "Complaint Share", "others": "Others",
            "ai_slide_title": "4. Gemini AI Root Cause Engineering Diagnosis",
        },
        "zh": {
            "deck_title": f"[B2B 定制竞标提案]\n{client_name} 遥控器品质与用户体验革新方案",
            "date": "编制日期", "target": "目标品牌",
            "s1_title": "1. 核心客户痛点 (VoC Pain Points)",
            "p_rank1": "🚨 第一缺陷指标:", "p_rank2": "⚠️ 第二缺陷指标:",
            "reviews_header": "💬 现场典型用户痛点原声:",
            "s2_title": "2. 我方 R&D 核心技术应对方案", "sol_for": "针对性应对方案:",
            "s3_title": "3. 高管汇报与商务竞标 Pitch",
            "exec_pitch": f"Executive Pitch for {client_name}",
            "chart_series": "负面占比", "others": "其他",
            "ai_slide_title": "4. Gemini AI 缺陷根因与工程诊断 (Root Cause)",
        },
        "vi": {
            "deck_title": f"[Đề Xuất Dự Án B2B]\nGiải pháp nâng cao chất lượng & UX điều khiển RCU {client_name}",
            "date": "Ngày lập", "target": "Thương hiệu mục tiêu",
            "s1_title": "1. Điểm đau cốt lõi của khách hàng (VoC)",
            "p_rank1": "🚨 Lỗi ưu tiên 1:", "p_rank2": "⚠️ Lỗi ưu tiên 2:",
            "reviews_header": "💬 Trích dẫn khiếu nại thực tế từ khách hàng:",
            "s2_title": "2. Giải pháp kỹ thuật R&D từ chúng tôi", "sol_for": "Giải pháp kỹ thuật cho",
            "s3_title": "3. Kịch bản Pitch đàm phán B2B & Báo cáo lãnh đạo",
            "exec_pitch": f"Executive Pitch for {client_name}",
            "chart_series": "Tỷ lệ lỗi", "others": "Khác",
            "ai_slide_title": "4. Chẩn Đoán Lỗi Gốc Rễ Bằng Gemini AI",
        },
        "th": {
            "deck_title": f"[ข้อเสนอโครงการ B2B]\nแนวทางยกระดับคุณภาพและ UX รีโมทคอนโทรลสำหรับ {client_name}",
            "date": "วันที่จัดทำ", "target": "แบรนด์เป้าหมาย",
            "s1_title": "1. ปัญหาสำคัญที่ลูกค้าพบ (VoC Pain Points)",
            "p_rank1": "🚨 ดัชนีข้อบกพร่องอันดับ 1:", "p_rank2": "⚠️ ดัชนีข้อบกพร่องอันดับ 2:",
            "reviews_header": "💬 ข้อความจริงจากลูกค้าที่ใช้งาน:",
            "s2_title": "2. โซลูชันทางเทคนิคจากทีม R&D ของเรา", "sol_for": "โซลูชันสำหรับ",
            "s3_title": "3. สคริปต์นำเสนอผู้บริหารและการเจรจา B2B",
            "exec_pitch": f"Executive Pitch for {client_name}",
            "chart_series": "สัดส่วนข้อร้องเรียน", "others": "อื่นๆ",
            "ai_slide_title": "4. การวินิจฉัยสาเหตุที่แท้จริงด้วย Gemini AI",
        },
    }
    L = t_titles.get(lang, t_titles["en"])

    # Slide 1: Cover
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(240, 244, 248)
    slide.shapes.title.text = L["deck_title"]
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
    slide.shapes.title.text_frame.paragraphs[0].font.bold = True
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 119, 180)
    today_date = datetime.now().strftime("%Y-%m-%d")
    date_lbl = L["date"]
    target_lbl = L["target"]
    slide.placeholders[1].text = f"{date_lbl}: {today_date}\n{target_lbl}: {client_name}"

    # Slide 2: VoC Pain Points + Native Vector Doughnut Chart
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    title_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.5), Inches(0.8))
    p_t = title_box.text_frame.paragraphs[0]
    p_t.text = L["s1_title"]
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(31, 119, 180)

    # Left Column: Text Pain Points
    left_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf2 = left_box.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = f"{L['p_rank1']} {p_issue}"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RGBColor(214, 39, 40)
    if s_issue:
        p = tf2.add_paragraph()
        p.text = f"{L['p_rank2']} {s_issue}"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 127, 14)
    p = tf2.add_paragraph()
    p.text = ""
    p = tf2.add_paragraph()
    p.text = L["reviews_header"]
    p.font.size = Pt(15)
    p.font.bold = True
    for r in reviews:
        p = tf2.add_paragraph()
        p.text = f'  - "{r}"'
        p.font.size = Pt(12)
        p.font.italic = True
        p.font.color.rgb = RGBColor(90, 90, 90)

    # Right Column: Native Vector Doughnut Chart
    if cat_counts is not None and not cat_counts.empty:
        chart_data = CategoryChartData()
        top_cats = cat_counts.head(3)
        other_cnt = cat_counts.iloc[3:].sum() if len(cat_counts) > 3 else 0

        cat_labels = [tr.CATEGORIES.get(k, {}).get(lang, k) for k in top_cats.index]
        cat_vals = list(top_cats.values)
        if other_cnt > 0:
            cat_labels.append(L["others"])
            cat_vals.append(other_cnt)

        chart_data.categories = cat_labels
        chart_data.add_series(L["chart_series"], cat_vals)

        chart_shape = slide2.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, Inches(5.4), Inches(1.5), Inches(4.2), Inches(4.8), chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(9)
        chart.legend.include_in_layout = False

        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(9)
        plot.data_labels.font.bold = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False

    # Slide 3: R&D Solutions
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = L["s2_title"]
    tf3 = slide3.placeholders[1].text_frame
    tf3.clear()
    p = tf3.add_paragraph()
    sol_lbl = L["sol_for"]
    p.text = f"✔️ [{p_issue}] {sol_lbl}"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(44, 160, 44)
    p = tf3.add_paragraph()
    p.text = f"  {p_sol}"
    p.font.size = Pt(15)
    if s_issue:
        p = tf3.add_paragraph()
        p.text = ""
        p = tf3.add_paragraph()
        p.text = f"✔️ [{s_issue}] {sol_lbl}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(44, 160, 44)
        p = tf3.add_paragraph()
        p.text = f"  {s_sol}"
        p.font.size = Pt(15)

    # Slide 4: Executive Pitch
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    slide4.shapes.title.text = L["s3_title"]
    sp = slide4.placeholders[1].element
    sp.getparent().remove(sp)
    txBox = slide4.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3.5))
    txBox.fill.solid()
    txBox.fill.fore_color.rgb = RGBColor(245, 245, 245)
    txBox.line.color.rgb = RGBColor(31, 119, 180)
    txBox.line.width = Pt(2)
    tf4 = txBox.text_frame
    tf4.word_wrap = True
    p = tf4.add_paragraph()
    p.text = L["exec_pitch"]
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 119, 180)
    p = tf4.add_paragraph()
    p.text = ""
    p = tf4.add_paragraph()
    p.text = f'"{pitch}"'
    p.font.size = Pt(20)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER

    # Optional Slide 5: Gemini AI Root Cause Analysis (If generated)
    if ai_analysis and not ai_analysis.startswith("⚠️"):
        slide5 = prs.slides.add_slide(prs.slide_layouts[1])
        slide5.shapes.title.text = L["ai_slide_title"]
        tf5 = slide5.placeholders[1].text_frame
        tf5.clear()
        tf5.word_wrap = True

        for line in ai_analysis.splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            p = tf5.add_paragraph()
            if line_str.startswith("### "):
                p.text = line_str.replace("### ", "").strip()
                p.font.size = Pt(16)
                p.font.bold = True
                p.font.color.rgb = RGBColor(31, 119, 180)
            elif line_str.startswith(("- ", "* ")):
                p.text = "  • " + line_str[2:].strip()
                p.font.size = Pt(12)
                p.font.color.rgb = RGBColor(60, 60, 60)
            else:
                p.text = line_str
                p.font.size = Pt(12)

    ppt_stream = io.BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)
    return ppt_stream

# -------------------------------------------------------------
# 3. 사이드바 다차원 필터 패널 & API Key 설정
# -------------------------------------------------------------
st.sidebar.header(txt("sidebar_filter"))

if st.sidebar.button(txt("sync_btn")):
    st.cache_data.clear()
    st.rerun()

st.sidebar.markdown("---")
st.sidebar.markdown(f"##### {txt('preset_title')}")

all_unique_brands = sorted(list(df["Brand"].unique()))
if "selected_brand_list" not in st.session_state:
    st.session_state["selected_brand_list"] = all_unique_brands

col_btn1, col_btn2 = st.sidebar.columns(2)
with col_btn1:
    if st.button(txt("btn_kr")):
        st.session_state["selected_brand_list"] = [
            b for b in KR_CHAMPIONS if b in all_unique_brands
        ]
        st.rerun()
    if st.button(txt("btn_na")):
        st.session_state["selected_brand_list"] = [
            b for b in NA_CHAMPIONS if b in all_unique_brands
        ]
        st.rerun()
    if st.button(txt("btn_africa")):
        st.session_state["selected_brand_list"] = [
            b for b in AFRICA_CHAMPIONS if b in all_unique_brands
        ]
        st.rerun()

with col_btn2:
    if st.button(txt("btn_eu")):
        st.session_state["selected_brand_list"] = [
            b for b in EU_CHAMPIONS if b in all_unique_brands
        ]
        st.rerun()
    if st.button(txt("btn_apac")):
        st.session_state["selected_brand_list"] = [
            b for b in APAC_CHAMPIONS if b in all_unique_brands
        ]
        st.rerun()
    if st.button(txt("btn_all")):
        st.session_state["selected_brand_list"] = all_unique_brands
        st.rerun()

st.sidebar.markdown("---")

device_types = sorted(list(df["Device_Type"].dropna().unique()))
selected_types = st.sidebar.multiselect(
    txt("filter_dev"),
    options=device_types,
    default=device_types,
    format_func=get_dev_name,
)

brand_subset = df[df["Device_Type"].isin(selected_types)]["Brand"].unique()
available_brands = sorted(list(brand_subset))
valid_defaults = [
    b for b in st.session_state["selected_brand_list"] if b in available_brands
]
if not valid_defaults:
    valid_defaults = available_brands

selected_brands = st.sidebar.multiselect(
    txt("filter_brand"), options=available_brands, default=valid_defaults
)

with st.sidebar.expander(txt("filter_hw"), expanded=False):
    all_powers = sorted(list(df["Power_Type"].unique()))
    sel_powers = st.multiselect(
        txt("filter_power"),
        all_powers,
        default=all_powers,
        format_func=get_power_name,
    )

    all_uis = sorted(list(df["UI_Type"].unique()))
    sel_uis = st.multiselect(
        txt("filter_ui"),
        all_uis,
        default=all_uis,
        format_func=get_ui_name,
    )

selected_regions = st.sidebar.multiselect(
    txt("filter_region"),
    options=sorted(df["Region"].unique()),
    default=sorted(df["Region"].unique()),
    format_func=get_region_name,
)

selected_cats = st.sidebar.multiselect(
    txt("filter_cat"),
    options=sorted(df["Category"].unique()),
    default=sorted(df["Category"].unique()),
    format_func=get_cat_name,
)

search_keyword = st.sidebar.text_input(
    txt("filter_kw"),
    placeholder=txt("kw_placeholder"),
)

# 사이드바 하단 Gemini API 키 입력기
with st.sidebar.expander("🤖 Gemini AI API Settings", expanded=False):
    saved_key = get_secret_api_key()
    user_api_key = st.text_input(
        txt("ai_key_input"),
        type="password",
        value=st.session_state.get("gemini_api_key", saved_key),
        placeholder="AIzaSy...",
        help=txt("ai_key_help"),
        key="gemini_key_input_box"
    )
    if user_api_key:
        st.session_state["gemini_api_key"] = user_api_key.strip()

filtered_df = df[
    (df["Device_Type"].isin(selected_types))
    & (df["Brand"].isin(selected_brands))
    & (df["Power_Type"].isin(sel_powers))
    & (df["UI_Type"].isin(sel_uis))
    & (df["Region"].isin(selected_regions))
    & (df["Category"].isin(selected_cats))
]

if search_keyword.strip():
    filtered_df = filtered_df[
        filtered_df["Review"]
        .astype(str)
        .str.contains(search_keyword.strip(), case=False, na=False)
    ]

# -------------------------------------------------------------
# 4. 상단 헤더 & 알림 배너
# -------------------------------------------------------------
st.title(txt("main_title"))
st.caption(txt("main_caption"))

if not filtered_df.empty:
    neg_only = filtered_df[filtered_df["Sentiment"] == "Negative"]
    top_3_cats = neg_only["Category"].value_counts().head(3)
    top_cat_names = [
        f"**{i+1}** {get_cat_name(cat)}" for i, cat in enumerate(top_3_cats.index)
    ]

    brand_neg_stats = (
        filtered_df.groupby("Brand")
        .apply(lambda x: (x["Sentiment"] == "Negative").mean())
        .reset_index(name="neg_rate")
    )
    brand_counts = filtered_df["Brand"].value_counts().reset_index()
    brand_counts.columns = ["Brand", "count"]
    brand_merged = pd.merge(brand_neg_stats, brand_counts, on="Brand")
    valid_brands = brand_merged[brand_merged["count"] >= 15].sort_values(
        by="neg_rate", ascending=False
    )
    highest_neg_brand = (
        valid_brands.iloc[0]["Brand"] if not valid_brands.empty else "-"
    )
    highest_neg_val = (
        f"{valid_brands.iloc[0]['neg_rate']*100:.1f}%"
        if not valid_brands.empty
        else "-"
    )

    b1, b2 = st.columns([2, 1])
    with b1:
        st.info(
            f"{txt('top3_alert')} {' | '.join(top_cat_names) if top_cat_names else '-'}"
        )
    with b2:
        st.warning(
            f"{txt('neg_alert')} **{highest_neg_brand}** ({txt('complaint_rate')} {highest_neg_val})"
        )

st.divider()

# -------------------------------------------------------------
# 5. 5대 탭 분할 구조
# -------------------------------------------------------------
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    txt("tab1"), txt("tab2"), txt("tab3"), txt("tab4"), txt("tab5")
])

# =============================================================
# TAB 1: 종합 모니터링 & 트렌드 + 인라인 드릴다운 (Drill-Down Inspector)
# =============================================================
with tab1:
    total_cnt = len(filtered_df)
    neg_cnt = len(filtered_df[filtered_df["Sentiment"] == "Negative"])
    neg_ratio = (neg_cnt / total_cnt * 100) if total_cnt > 0 else 0
    top_cat_mode = filtered_df["Category"].mode()
    top_cat = get_cat_name(top_cat_mode.iloc[0]) if not top_cat_mode.empty else "-"

    m1, m2, m3, m4 = st.columns(4)
    m1.metric(txt("m_total"), f"{total_cnt:,}")
    m2.metric(
        txt("m_neg_rate"),
        f"{neg_ratio:.1f}%",
        delta=f"{neg_cnt:,}",
        delta_color="inverse",
    )
    m3.metric(txt("m_top_issue"), top_cat)
    m4.metric(txt("m_brand_cnt"), f"{filtered_df['Brand'].nunique()}")

    st.write("")
    col_t1, col_t2 = st.columns(2)
    selected_slice = None

    with col_t1:
        st.subheader(txt("chart_cat_dist"))
        if not filtered_df.empty:
            cat_counts = filtered_df["Category"].value_counts().reset_index()
            cat_counts.columns = ["Category", "Count"]
            cat_counts["Category_Label"] = cat_counts["Category"].apply(get_cat_name)

            # 호버 툴팁용 정의 및 주요 사례 주입 (이미지와 동일한 📌 정의 & 💬 주요 사례 서식)
            cat_counts["Description"] = cat_counts["Category"].map(lambda c: tr.get_cat_desc(c, curr_lang))
            cat_counts["Examples"] = cat_counts["Category"].map(lambda c: tr.get_cat_example(c, curr_lang))

            customdata = cat_counts[["Description", "Examples"]].values

            hovertemplate = (
                "<b>%{label}</b><br>"
                + f"{txt('share_label')}: <b>%{{percent}}</b> (%{{value:,}}{txt('count_unit')})<br><br>"
                + f"📌 <b>{txt('lbl_definition')}:</b> %{{customdata[0]}}<br>"
                + f"💬 <b>{txt('lbl_key_examples')}:</b> %{{customdata[1]}}"
                + "<extra></extra>"
            )

            fig1 = go.Figure(
                data=[
                    go.Pie(
                        labels=cat_counts["Category_Label"],
                        values=cat_counts["Count"],
                        hole=0.4,
                        textinfo="percent",
                        customdata=customdata,
                        hovertemplate=hovertemplate,
                        marker=dict(colors=px.colors.qualitative.Set3),
                    )
                ]
            )
            fig1.update_layout(
                margin=dict(t=20, b=20, l=10, r=10),
                legend=dict(orientation="v", x=1.05, y=0.5),
            )

            # 차트 클릭 지원 (Streamlit 1.35 이상 동적 감지)
            try:
                import inspect
                if "on_select" in inspect.signature(st.plotly_chart).parameters:
                    chart_event = st.plotly_chart(
                        fig1,
                        use_container_width=True,
                        on_select="rerun",
                        selection_mode="points",
                        key="pie_chart_cat_dist",
                    )
                    if (
                        chart_event
                        and isinstance(chart_event, dict)
                        and chart_event.get("selection")
                        and chart_event["selection"].get("points")
                    ):
                        p_idx = chart_event["selection"]["points"][0].get("point_index")
                        if p_idx is not None and 0 <= p_idx < len(cat_counts):
                            selected_slice = cat_counts.iloc[p_idx]["Category"]
                else:
                    st.plotly_chart(fig1, use_container_width=True)
            except Exception:
                st.plotly_chart(fig1, use_container_width=True)

    with col_t2:
        st.subheader(txt("chart_dev_dist"))
        if not filtered_df.empty:
            dev_df = filtered_df.copy()
            dev_df["Device_Type_Label"] = dev_df["Device_Type"].apply(get_dev_name)
            dev_df["Category_Label"] = dev_df["Category"].apply(get_cat_name)
            dev_cat = (
                dev_df.groupby(["Device_Type_Label", "Category_Label"])
                .size()
                .reset_index(name="건수")
            )
            fig2 = px.bar(
                dev_cat,
                x="Device_Type_Label",
                y="건수",
                color="Category_Label",
                barmode="stack",
                color_discrete_sequence=px.colors.qualitative.Set3,
                labels={
                    "Device_Type_Label": txt("filter_dev"),
                    "건수": txt("count_unit"),
                    "Category_Label": txt("filter_cat"),
                },
            )
            fig2.update_layout(margin=dict(t=20, b=20, l=10, r=10))
            st.plotly_chart(fig2, use_container_width=True)

    # -------------------------------------------------------------
    # 4번 연계: 인라인 심층 드릴다운 & 현장 VoC 원문 5선 (Drill-Down Inspector)
    # -------------------------------------------------------------
    st.write("")
    st.markdown("---")
    st.subheader(txt("drilldown_title"))
    st.caption(txt("drilldown_caption"))

    if not filtered_df.empty:
        available_inspect_cats = sorted(list(filtered_df["Category"].unique()))

        # 차트 클릭 선택 또는 직전 선택 상태 동기화
        if selected_slice and selected_slice in available_inspect_cats:
            default_inspect_cat = selected_slice
        elif "chosen_drilldown_cat" in st.session_state and st.session_state["chosen_drilldown_cat"] in available_inspect_cats:
            default_inspect_cat = st.session_state["chosen_drilldown_cat"]
        else:
            default_inspect_cat = (
                top_cat_mode.iloc[0]
                if (not top_cat_mode.empty and top_cat_mode.iloc[0] in available_inspect_cats)
                else available_inspect_cats[0]
            )

        inspect_cat = st.selectbox(
            txt("select_inspect_cat"),
            options=available_inspect_cats,
            index=available_inspect_cats.index(default_inspect_cat),
            format_func=get_cat_name,
            key="inline_inspect_cat_select",
        )
        st.session_state["chosen_drilldown_cat"] = inspect_cat

        col_dd1, col_dd2 = st.columns([1.1, 1.3])
        with col_dd1:
            st.markdown(f"##### {txt('cat_summary_card')}")
            cat_desc = tr.get_cat_desc(inspect_cat, curr_lang)
            cat_ex = tr.get_cat_example(inspect_cat, curr_lang)
            cat_sol = get_action_sol(inspect_cat)
            cat_count_val = len(filtered_df[filtered_df["Category"] == inspect_cat])
            cat_ratio_val = (cat_count_val / total_cnt * 100) if total_cnt > 0 else 0

            st.markdown(
                f"""
                <div style="background-color: #f8f9fa; padding: 18px; border-radius: 10px; border-left: 5px solid #1f77b4; margin-bottom: 12px; box-shadow: 0 1px 3px rgba(0,0,0,0.05);">
                    <div style="font-size: 1.1em; font-weight: bold; color: #1f77b4; margin-bottom: 6px;">🏷️ {get_cat_name(inspect_cat)} <span style="font-size: 0.85em; color: #666;">({cat_count_val:,}{txt('count_unit')} · {cat_ratio_val:.1f}%)</span></div>
                    <div style="font-size: 0.92em; line-height: 1.5; margin-bottom: 8px;">📌 <b>{txt('lbl_definition')}:</b> {cat_desc}</div>
                    <div style="font-size: 0.92em; line-height: 1.5; color: #444;">💬 <b>{txt('lbl_key_examples')}:</b> {cat_ex}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.success(f"**🛠️ {txt('hw_fw_proposal')}**\n\n{cat_sol}")

        with col_dd2:
            st.markdown(f"##### {txt('cat_reviews_card')}")
            cat_reviews = filtered_df[
                (filtered_df["Category"] == inspect_cat)
                & (filtered_df["Sentiment"] == "Negative")
            ]["Review"].dropna()

            if cat_reviews.empty:
                cat_reviews = filtered_df[filtered_df["Category"] == inspect_cat]["Review"].dropna()

            if not cat_reviews.empty:
                sample_5_reviews = cat_reviews.head(5).tolist()
                for idx, rev_text in enumerate(sample_5_reviews):
                    st.markdown(
                        f"""
                        <div style="background-color: #ffffff; padding: 12px 16px; border-radius: 8px; border: 1px solid #e2e8f0; border-left: 4px solid #ef4444; margin-bottom: 8px; font-size: 0.92em; line-height: 1.5;">
                            <span style="color: #ef4444; font-weight: bold; margin-right: 6px;">#{idx+1}</span> "{rev_text}"
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
            else:
                st.info(txt("no_cat_reviews"))

    st.write("")
    st.markdown("---")
    st.subheader(txt("chart_trend"))
    if not filtered_df.empty:
        trend_df = filtered_df.copy()
        trend_df["YearWeek"] = trend_df["Parsed_Date"].dt.strftime("%Y-W%W")
        trend_df["Category_Label"] = trend_df["Category"].apply(get_cat_name)
        weekly_trend = (
            trend_df.groupby(["YearWeek", "Category_Label"])
            .size()
            .reset_index(name="건수")
        )
        fig_trend = px.line(
            weekly_trend,
            x="YearWeek",
            y="건수",
            color="Category_Label",
            markers=True,
            color_discrete_sequence=px.colors.qualitative.Set3,
            labels={
                "YearWeek": txt("week_unit"),
                "건수": txt("count_unit"),
                "Category_Label": txt("filter_cat"),
            },
        )
        fig_trend.update_layout(margin=dict(t=20, b=20, l=10, r=10))
        st.plotly_chart(fig_trend, use_container_width=True)

# =============================================================
# TAB 2: 브랜드 1:1 벤치마크
# =============================================================
with tab2:
    st.subheader(txt("t2_title"))
    st.caption(txt("t2_caption"))

    all_b = sorted(list(df["Brand"].dropna().unique()))
    if len(all_b) >= 2:
        c_b1, c_b2 = st.columns(2)
        with c_b1:
            brand_a = st.selectbox(
                txt("brand_a_label"),
                all_b,
                index=all_b.index("Samsung") if "Samsung" in all_b else 0,
            )
        with c_b2:
            cand_b = [b for b in all_b if b != brand_a]
            brand_b = st.selectbox(
                txt("brand_b_label"),
                cand_b,
                index=(
                    cand_b.index("MultiChoice (DStv)")
                    if "MultiChoice (DStv)" in cand_b
                    else 0
                ),
            )

        df_a, df_b = df[df["Brand"] == brand_a], df[df["Brand"] == brand_b]
        a_cnt, b_cnt = len(df_a), len(df_b)
        a_neg_rate = (
            (len(df_a[df_a["Sentiment"] == "Negative"]) / a_cnt * 100)
            if a_cnt > 0
            else 0
        )
        b_neg_rate = (
            (len(df_b[df_b["Sentiment"] == "Negative"]) / b_cnt * 100)
            if b_cnt > 0
            else 0
        )

        k1, k2, km, k3, k4 = st.columns([2, 2, 1, 2, 2])
        k1.metric(f"[{brand_a}] {txt('m_total')}", f"{a_cnt:,}")
        k2.metric(f"[{brand_a}] {txt('m_neg_rate')}", f"{a_neg_rate:.1f}%")
        km.markdown(
            "<h3 style='text-align: center; color: gray;'>VS</h3>",
            unsafe_allow_html=True,
        )
        k3.metric(f"[{brand_b}] {txt('m_total')}", f"{b_cnt:,}")
        k4.metric(
            f"[{brand_b}] {txt('m_neg_rate')}",
            f"{b_neg_rate:.1f}%",
            delta=f"{b_neg_rate - a_neg_rate:+.1f}%p",
            delta_color="inverse",
        )

        df_ab = df[df["Brand"].isin([brand_a, brand_b])].copy()
        if not df_ab.empty:
            df_ab["Category_Label"] = df_ab["Category"].apply(get_cat_name)
            cat_ab = (
                df_ab.groupby(["Brand", "Category_Label"])
                .size()
                .reset_index(name="건수")
            )
            tot_ab = cat_ab.groupby("Brand")["건수"].transform("sum")
            cat_ab["비중(%)"] = (cat_ab["건수"] / tot_ab * 100).round(1)

            fig_comp = px.bar(
                cat_ab,
                x="Category_Label",
                y="비중(%)",
                color="Brand",
                barmode="group",
                color_discrete_sequence=["#1f77b4", "#ff7f0e"],
                labels={
                    "Category_Label": txt("filter_cat"),
                    "비중(%)": txt("ratio_unit"),
                },
            )
            fig_comp.update_layout(xaxis_tickangle=-30, margin=dict(t=20, b=20))
            st.plotly_chart(fig_comp, use_container_width=True)

        st.markdown(f"#### {txt('rd_action_title')}")
        col_act_a, col_act_b = st.columns(2)
        with col_act_a:
            top_issue_a = df_a[df_a["Sentiment"] == "Negative"][
                "Category"
            ].mode()
            act_cat_a = (
                top_issue_a.iloc[0] if not top_issue_a.empty else "기타/일반"
            )
            rec_a = get_action_sol(act_cat_a)
            st.success(f"**[{brand_a}] {txt('core_weakness')}: {get_cat_name(act_cat_a)}**\n\n{txt('hw_fw_proposal')}\n- {rec_a}")

        with col_act_b:
            top_issue_b = df_b[df_b["Sentiment"] == "Negative"][
                "Category"
            ].mode()
            act_cat_b = (
                top_issue_b.iloc[0] if not top_issue_b.empty else "기타/일반"
            )
            rec_b = get_action_sol(act_cat_b)
            st.warning(f"**[{brand_b}] {txt('core_weakness')}: {get_cat_name(act_cat_b)}**\n\n{txt('hw_fw_proposal')}\n- {rec_b}")

# =============================================================
# TAB 3: 키워드 & 연관구
# =============================================================
with tab3:
    st.subheader(txt("t3_title"))
    col_t3_1, col_t3_2 = st.columns(2)
    neg_reviews = filtered_df[filtered_df["Sentiment"] == "Negative"][
        "Review"
    ].astype(str)
    stopwords = {
        "the", "and", "for", "with", "that", "this", "not", "but", "are",
        "was", "were", "remote", "control", "device", "tv", "app", "work",
        "works", "one", "all", "you", "can", "cant", "just", "very",
        "also", "from", "have", "리모컨", "작동", "사용", "제품", "진짜",
        "너무", "그냥", "다시",
    }
    single_words, bigram_phrases = [], []
    for rev in neg_reviews:
        tokens = [
            w
            for w in re.findall(r"[a-zA-Z가-힣]{2,}", rev.lower())
            if w not in stopwords and len(w) > 2
        ]
        single_words.extend(tokens)
        for i in range(len(tokens) - 1):
            bigram_phrases.append(f"{tokens[i]} {tokens[i+1]}")

    with col_t3_1:
        st.markdown(f"##### {txt('top_keywords')}")
        if single_words:
            df_words = pd.DataFrame(
                Counter(single_words).most_common(10),
                columns=["Keyword", "Frequency"],
            )
            fig_words = px.bar(
                df_words,
                x="Frequency",
                y="Keyword",
                orientation="h",
                color="Frequency",
                color_continuous_scale="Reds",
            )
            fig_words.update_layout(
                yaxis={"autorange": "reversed"},
                margin=dict(t=20, b=20, l=10, r=10),
            )
            st.plotly_chart(fig_words, use_container_width=True)

    with col_t3_2:
        st.markdown(f"##### {txt('top_bigrams')}")
        if bigram_phrases:
            df_bigrams = pd.DataFrame(
                Counter(bigram_phrases).most_common(10),
                columns=["Bigram", "Frequency"],
            )
            fig_bigrams = px.bar(
                df_bigrams,
                x="Frequency",
                y="Bigram",
                orientation="h",
                color="Frequency",
                color_continuous_scale="Oranges",
            )
            fig_bigrams.update_layout(
                yaxis={"autorange": "reversed"},
                margin=dict(t=20, b=20, l=10, r=10),
            )
            st.plotly_chart(fig_bigrams, use_container_width=True)

# =============================================================
# TAB 4: 원문 DB & 다운로드 (다국어 테이블 및 컬럼 번역 적용)
# =============================================================
with tab4:
    st.subheader(txt("t4_title"))
    display_cols = [
        c
        for c in [
            "Device_Type", "Power_Type", "UI_Type", "Region", "Brand",
            "Category", "Sentiment", "Source", "Review", "Collected_Date",
        ]
        if c in filtered_df.columns
    ]

    table_df = filtered_df[display_cols].copy()
    if "Device_Type" in table_df.columns:
        table_df["Device_Type"] = table_df["Device_Type"].apply(get_dev_name)
    if "Power_Type" in table_df.columns:
        table_df["Power_Type"] = table_df["Power_Type"].apply(get_power_name)
    if "UI_Type" in table_df.columns:
        table_df["UI_Type"] = table_df["UI_Type"].apply(get_ui_name)
    if "Region" in table_df.columns:
        table_df["Region"] = table_df["Region"].apply(get_region_name)
    if "Category" in table_df.columns:
        table_df["Category"] = table_df["Category"].apply(get_cat_name)
    if "Sentiment" in table_df.columns:
        table_df["Sentiment"] = table_df["Sentiment"].apply(get_sentiment_name)

    table_df = table_df.rename(columns={c: get_col_name(c) for c in display_cols})

    st.dataframe(table_df, use_container_width=True, height=500)
    today_date = datetime.now().strftime("%Y%m%d")

    csv_bytes = table_df.to_csv(index=False).encode("utf-8-sig")
    st.download_button(
        label=txt("btn_csv_down"),
        data=csv_bytes,
        file_name=f"rcu_voc_{curr_lang}_{today_date}.csv",
        mime="text/csv",
    )

# =============================================================
# TAB 5: B2B 수주 & 임원 보고용 Pitch (Gemini AI 심층 진단 탑재)
# =============================================================
with tab5:
    st.subheader(txt("t5_title"))
    st.caption(txt("t5_caption"))

    all_pitch_brands = sorted(list(df["Brand"].dropna().unique()))
    target_client = st.selectbox(
        txt("target_client_label"),
        all_pitch_brands,
        index=(
            all_pitch_brands.index("KT (지니TV)")
            if "KT (지니TV)" in all_pitch_brands
            else 0
        ),
    )
    st.divider()

    client_df = df[df["Brand"] == target_client]
    client_neg = client_df[client_df["Sentiment"] == "Negative"]

    if client_neg.empty:
        st.info(txt("no_neg_voc", client=target_client))
    else:
        client_cat_counts = client_neg["Category"].value_counts()
        top_issues = client_cat_counts.head(2)
        primary_issue_raw = top_issues.index[0]
        secondary_issue_raw = (
            top_issues.index[1] if len(top_issues) > 1 else None
        )

        primary_issue = get_cat_name(primary_issue_raw)
        secondary_issue = (
            get_cat_name(secondary_issue_raw) if secondary_issue_raw else None
        )

        col_p1, col_p2 = st.columns([1, 1.4])
        with col_p1:
            st.markdown(f"#### {txt('pain_point_title', client=target_client)}")
            st.error(
                f"**{txt('rank1_defect')}:** {primary_issue} ({txt('complaints_cnt', cnt=top_issues.iloc[0])})"
            )
            if secondary_issue:
                st.warning(
                    f"**{txt('rank2_defect')}:** {secondary_issue} ({txt('complaints_cnt', cnt=top_issues.iloc[1])})"
                )

            st.markdown(f"{txt('sample_reviews_title')}")
            valid_reviews = client_neg[
                ~client_neg["Review"].astype(str).str.contains(NOISE_FILTER_REGEX)
            ]["Review"]
            if valid_reviews.empty:
                valid_reviews = client_neg["Review"]
            sample_reviews = valid_reviews.dropna().head(3).tolist()
            for r in sample_reviews:
                st.caption(f"- {r}")

        with col_p2:
            st.markdown(f"#### {txt('rd_solution_title')}")
            sol_1 = get_action_sol(primary_issue_raw)
            st.success(f"✔️ **{txt('counter_sol', cat=primary_issue)}**\n\n{sol_1}")
            sol_2 = ""
            if secondary_issue_raw:
                sol_2 = get_action_sol(secondary_issue_raw)
                st.info(f"✔️ **{txt('counter_sol', cat=secondary_issue)}**\n\n{sol_2}")

        st.write("")
        st.markdown("---")
        st.markdown(f"#### {txt('pitch_title')}")

        dev_type = get_dev_name(
            BRAND_DIRECTORY.get(target_client, {}).get("type", "디바이스")
        )

        pitch_templates = {
            "ko": (
                f"현재 {target_client} 가입자들이 가장 큰 불편을 호소하고 있는"
                f" '{primary_issue}' 결함을 당사의 {dev_type} 특화 차세대 아키텍처로"
                " 완벽하게 해결하겠습니다. 단순한 납품 단가 절감이 아닌, 반품 및 CS"
                f" 콜센터 인입률을 원천 차단하여 {target_client}의 브랜드 신뢰도 제고와"
                " 운영 비용(OPEX)의 획기적 절감을 제안합니다."
            ),
            "en": (
                f"We will resolve the critical '{primary_issue}' defect"
                f" reported by {target_client} users with our next-generation"
                f" architecture tailored for {dev_type}. Beyond unit cost"
                " reductions, we propose eliminating return rates and customer"
                " support calls at the root, enhancing brand trust and"
                " drastically reducing OPEX."
            ),
            "zh": (
                f"我们将通过专为 {dev_type} 定制的下一代高可靠架构，彻底解决"
                f" {target_client} 用户当前反馈最强烈的 '{primary_issue}'"
                " 缺陷。不仅仅是降低元器件采购成本，更能从源头杜绝退货与售后呼叫中心成本，提升品牌美誉度并大幅降低运营支出"
                " (OPEX)。"
            ),
            "vi": (
                f"Chúng tôi sẽ giải quyết dứt điểm lỗi '{primary_issue}' mà"
                f" người dùng {target_client} đang phàn nàn nhiều nhất bằng kiến"
                f" trúc thế hệ mới dành riêng cho {dev_type}. Không chỉ giảm"
                " giá thành linh kiện, chúng tôi cam kết loại bỏ tỷ lệ đổi trả"
                " và cuộc gọi hỗ trợ kỹ thuật, nâng cao uy tín thương hiệu và"
                " tiết kiệm chi phí vận hành (OPEX)."
            ),
            "th": (
                f"เราพร้อมแก้ไขปัญหา '{primary_issue}' ที่ผู้ใช้งาน"
                f" {target_client} ร้องเรียนมากที่สุดด้วยสถาปัตยกรรมรุ่นใหม่ที่ออกแบบมาโดยเฉพาะสำหรับ"
                f" {dev_type} นอกเหนือจากการลดต้นทุนชิ้นส่วนแล้ว"
                " เรายังช่วยลดอัตราการเคลมสินค้าและภาระศูนย์บริการลูกค้าได้อย่างสิ้นเชิง"
                " เสริมสร้างความเชื่อมั่นต่อแบรนด์และลดค่าใช้จ่ายในการดำเนินงาน"
                " (OPEX) อย่างยั่งยืน"
            ),
        }
        final_pitch = pitch_templates.get(curr_lang, pitch_templates["en"])

        st.markdown(
            f"""
            <div style="background-color: #f8f9fa; padding: 22px; border-radius: 10px; border-left: 6px solid #1f77b4; box-shadow: 0 2px 4px rgba(0,0,0,0.05);">
                <span style="font-size: 0.9em; font-weight: bold; color: #1f77b4; text-transform: uppercase;">Executive Pitch for {target_client} ({curr_lang.upper()})</span>
                <p style="font-size: 1.15em; line-height: 1.65; color: #212529; margin-top: 8px; font-weight: 500;">
                    "{final_pitch}"
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.write("")
        st.markdown("---")

        # ---------------------------------------------------------
        # Gemini AI VoC 심층 진단 & 결함 원인(Root Cause) 섹션
        # ---------------------------------------------------------
        st.subheader(txt("ai_section_title"))
        ai_session_key = f"ai_insight_{target_client}_{curr_lang}"

        active_api_key = st.session_state.get(
            "gemini_api_key",
            get_secret_api_key()
        )

        col_ai_btn, col_ai_info = st.columns([1, 2.5])
        with col_ai_btn:
            trigger_ai = st.button(txt("ai_btn"), type="primary")

        if trigger_ai:
            if not active_api_key:
                st.warning(txt("ai_key_missing"))
            else:
                with st.spinner(txt("ai_loading")):
                    ai_result = analyze_voc_with_gemini(
                        active_api_key,
                        target_client,
                        primary_issue,
                        secondary_issue,
                        sample_reviews,
                        lang=curr_lang
                    )
                    st.session_state[ai_session_key] = ai_result

        # AI 결과 출력 컨테이너
        current_ai_text = st.session_state.get(ai_session_key, None)
        if current_ai_text:
            st.markdown(
                f"""
                <div style="background-color: #f0f7ff; padding: 24px; border-radius: 12px; border: 1.5px solid #90caf9; margin-top: 15px; margin-bottom: 20px;">
                """,
                unsafe_allow_html=True
            )
            st.markdown(current_ai_text)
            st.markdown("</div>", unsafe_allow_html=True)

        st.write("")
        ppt_file = create_pptx_deck(
            target_client,
            primary_issue,
            secondary_issue,
            sol_1,
            sol_2,
            final_pitch,
            sample_reviews,
            cat_counts=client_cat_counts,
            ai_analysis=current_ai_text,
            lang=curr_lang,
        )
        st.download_button(
            label=txt("btn_pptx_down", client=target_client),
            data=ppt_file,
            file_name=f"B2B_Pitch_{target_client.replace('/', '_')}_{curr_lang}_{today_date}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
